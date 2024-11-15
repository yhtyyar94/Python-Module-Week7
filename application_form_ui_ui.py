from PyQt6.QtWidgets import (
    QApplication,
    QWidget,
    QLabel,
    QLineEdit,
    QPushButton,
    QVBoxLayout,
    QComboBox,
    QRadioButton,
    QButtonGroup,
    QCheckBox,
    QMessageBox,
    QTextEdit,
    QScrollArea,
)
from PyQt6.QtCore import Qt
import sys
from backend.send_email import send_email
from backend.list_files import list_drive_files
from backend.download_file import download_file
from backend.upload_file import update_file
from openpyxl import load_workbook
import datetime
from db_controllers.connect import connect


class VITForm(QWidget):
    def __init__(self):
        super().__init__()
        self.setupUi()

    def setupUi(self):
        self.setFixedSize(800, 600)
        self.setWindowTitle("VIT Projesi Başvuru Formu")

        self.form_layout = QVBoxLayout()

        def create_label(text):
            label = QLabel(text)
            label.setStyleSheet("font-weight: bold;")
            return label

        # Email field
        self.email_label = create_label("Email Adresiniz:")
        self.email_input = QLineEdit()
        self.form_layout.addWidget(self.email_label)
        self.form_layout.addWidget(self.email_input)

        # Adınız Soyadınız
        self.name_label = create_label("Adınız Soyadınız:")
        self.name_input = QLineEdit()
        self.form_layout.addWidget(self.name_label)
        self.form_layout.addWidget(self.name_input)

        # Telefon Numaranız
        self.phone_label = create_label("Telefon Numaranız:")
        self.phone_input = QLineEdit()
        self.form_layout.addWidget(self.phone_label)
        self.form_layout.addWidget(self.phone_input)

        # Posta Kodu
        self.postcode_label = create_label("Posta Kodunuz:")
        self.postcode_input = QLineEdit()
        self.form_layout.addWidget(self.postcode_label)
        self.form_layout.addWidget(self.postcode_input)

        # Yaşadığınız Eyalet
        self.state_label = create_label("Yaşadığınız Eyalet:")
        self.state_combo = QComboBox()
        self.state_combo.setPlaceholderText("Eyalet Seçiniz")
        self.state_combo.addItems(
            [
                "Drenthe",
                "Flevoland",
                "Friesland",
                "Gelderland",
                "Groningen",
                "Limburg",
                "Noord-Brabant",
                "Noord-Holland",
                "Overijssel",
                "Utrecht",
                "Zeeland",
                "Zuid-Holland",
            ]
        )
        self.form_layout.addWidget(self.state_label)
        self.form_layout.addWidget(self.state_combo)

        # Ekonomik Durum (RadioButton)
        self.economic_status_label = create_label("Ekonomik Durumunuz:")
        self.economic_status_group = QButtonGroup(self)
        self.uitkering_btn = QRadioButton("Uitkering Alıyorum")
        self.part_time_btn = QRadioButton(
            "Kısmen Uitkering Alıyorum - Parttime Çalışıyorum"
        )
        self.uwv_btn = QRadioButton("UWV den uitkering alıyorum")
        self.no_uitkering_btn = QRadioButton("Çalışıyorum (Uitkerinden Çıktım)")
        self.no_income_btn = QRadioButton("Kampta kaldığım için uitkering almıyorum")
        self.form_layout.addWidget(self.economic_status_label)
        self.form_layout.addWidget(self.uitkering_btn)
        self.form_layout.addWidget(self.part_time_btn)
        self.form_layout.addWidget(self.uwv_btn)
        self.form_layout.addWidget(self.no_uitkering_btn)
        self.form_layout.addWidget(self.no_income_btn)
        self.economic_status_group.addButton(self.uitkering_btn)
        self.economic_status_group.addButton(self.part_time_btn)
        self.economic_status_group.addButton(self.uwv_btn)
        self.economic_status_group.addButton(self.no_uitkering_btn)
        self.economic_status_group.addButton(self.no_income_btn)

        # Şu anda bir dil kursuna devam ediyor musunuz?
        self.course_label = create_label(
            "Şu anda bir dil kursuna devam ediyor musunuz?"
        )
        self.course_yes = QRadioButton("Evet")
        self.course_no = QRadioButton("Hayır")
        self.course_group = QButtonGroup(self)
        self.course_group.addButton(self.course_yes)
        self.course_group.addButton(self.course_no)
        self.form_layout.addWidget(self.course_label)
        self.form_layout.addWidget(self.course_yes)
        self.form_layout.addWidget(self.course_no)

        # Yabancı Dil Seviyeniz (Checkbox)
        self.language_label = create_label("Yabancı Dil Seviyeniz:")
        self.a1_checkbox = QCheckBox("A1")
        self.a2_checkbox = QCheckBox("A2")
        self.b1_checkbox = QCheckBox("B1")
        self.b2_checkbox = QCheckBox("B2")
        self.form_layout.addWidget(self.language_label)
        self.form_layout.addWidget(self.a1_checkbox)
        self.form_layout.addWidget(self.a2_checkbox)
        self.form_layout.addWidget(self.b1_checkbox)
        self.form_layout.addWidget(self.b2_checkbox)

        # IT Kursu/Tecrübe/Projeler (Multiple Text Areas)
        self.it_course_label = create_label(
            "Başka bir IT kursu aldınız mı? (Detayları yazın)"
        )
        self.it_course_input = QTextEdit()
        self.form_layout.addWidget(self.it_course_label)
        self.form_layout.addWidget(self.it_course_input)

        self.it_experience_label = create_label(
            "Daha önce herhangi bir IT iş tecrübeniz var mı?"
        )
        self.it_experience_input = QTextEdit()
        self.form_layout.addWidget(self.it_experience_label)
        self.form_layout.addWidget(self.it_experience_input)

        self.project_label = create_label(
            "Öğretmenlik projesi veya Leerwerktraject projesine dahil misiniz?"
        )
        self.project_input = QTextEdit()
        self.form_layout.addWidget(self.project_label)
        self.form_layout.addWidget(self.project_input)

        # IT Sektörü (CheckBoxes)
        self.it_sector_label = create_label(
            "IT sektöründe hangi bölümde çalışmak istersiniz?"
        )
        self.cloud_checkbox = QCheckBox("Cloud Engineer")
        self.tester_checkbox = QCheckBox("Tester")
        self.scrum_checkbox = QCheckBox("Scrum Master")
        self.devops_checkbox = QCheckBox("DevOps Engineer")
        self.web_checkbox = QCheckBox("Web Tasarım")
        self.software_checkbox = QCheckBox("Software Developer")
        self.data_checkbox = QCheckBox("Data Engineer")
        self.cyber_checkbox = QCheckBox("Cyber Security")
        self.form_layout.addWidget(self.it_sector_label)
        self.form_layout.addWidget(self.cloud_checkbox)
        self.form_layout.addWidget(self.tester_checkbox)
        self.form_layout.addWidget(self.scrum_checkbox)
        self.form_layout.addWidget(self.devops_checkbox)
        self.form_layout.addWidget(self.web_checkbox)
        self.form_layout.addWidget(self.software_checkbox)
        self.form_layout.addWidget(self.data_checkbox)
        self.form_layout.addWidget(self.cyber_checkbox)

        # IT Motivasyon (Text Area)
        self.motivation_label = create_label(
            "IT sektöründe kariyer yapmak için sizi harekete geçiren motivasyon nedir?"
        )
        self.motivation_input = QTextEdit()
        self.form_layout.addWidget(self.motivation_label)
        self.form_layout.addWidget(self.motivation_input)

        # Submit Button
        self.submit_button = QPushButton("Gönder")
        self.submit_button.clicked.connect(self.confirm_submission)
        self.form_layout.addWidget(self.submit_button)

        # Scroll Area
        scroll_area = QScrollArea()
        scroll_area.setWidgetResizable(True)
        scroll_content = QWidget()
        scroll_content.setLayout(self.form_layout)
        scroll_area.setWidget(scroll_content)

        # Set main layout
        main_layout = QVBoxLayout(self)
        main_layout.addWidget(scroll_area)
        self.setLayout(main_layout)

    def confirm_submission(self):
        form_data = [
            datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            self.name_input.text().strip() or "",
            self.email_input.text().strip() or "",
            self.phone_input.text().strip() or "",
            self.postcode_input.text().strip() or "",
            self.state_combo.currentText().strip() or "",
            (
                self.economic_status_group.checkedButton().text()
                if self.economic_status_group.checkedButton()
                else ""
            ),
            (
                "Evet"
                if self.course_yes.isChecked()
                else ("Hayır" if self.course_no.isChecked() else "")
            ),
            ", ".join(
                [
                    checkbox.text()
                    for checkbox in [
                        self.a1_checkbox,
                        self.a2_checkbox,
                        self.b1_checkbox,
                        self.b2_checkbox,
                    ]
                    if checkbox.isChecked()
                ]
            )
            or "",
            self.it_course_input.toPlainText().strip() or "",
            self.it_experience_input.toPlainText().strip() or "",
            self.project_input.toPlainText().strip() or "",
            ", ".join(
                [
                    checkbox.text()
                    for checkbox in [
                        self.cloud_checkbox,
                        self.tester_checkbox,
                        self.scrum_checkbox,
                        self.devops_checkbox,
                        self.web_checkbox,
                        self.software_checkbox,
                        self.data_checkbox,
                        self.cyber_checkbox,
                    ]
                    if checkbox.isChecked()
                ]
            )
            or "",
            self.motivation_input.toPlainText().strip() or "",
            "ATANMADI",
            "VIT5",
        ]
        self.submit_form(form_data)

    def submit_form(self, form_data):
        try:
            email_body = f"""
            Merhabalar,

            Aşağıdaki bilgilerle Vit projesine kayıt oldunuz:

            Zaman Damgası: {form_data[0]}
            Adınız Soyadınız: {form_data[1]}
            Email Adresiniz: {form_data[2]}
            Telefon Numaranız: {form_data[3]}
            Posta Kodunuz: {form_data[4]}
            Yaşadığınız Eyalet: {form_data[5]}
            Ekonomik Durumunuz: {form_data[6]}
            Dil Kursuna Devam: {form_data[7]}
            Dil Seviyesi: {form_data[8]}
            IT Kursu: {form_data[9]}
            IT Deneyimi: {form_data[10]}
            Katıldığı Proje: {form_data[11]}
            Tercih Edilen Sektör: {form_data[12]}
            Motivasyon: {form_data[13]}
            """

            send_email(form_data[2], email_body, "VIT Projesi Başvurunuz Hk.")
            send_email("vit5crmproject@gmail.com", email_body, "Yeni Başvuru")

            # send the data to database
            conn = connect("crm")
            cur = conn.cursor()
            cur.execute("Select * from applications LIMIT 1")
            headers = [desc[0] for desc in cur.description]
            columns = ", ".join([f'"{header}"' for header in headers[1:]])
            placeholders = ", ".join(["%s"] * len(headers[1:]))
            insert_query = (
                f"INSERT INTO applications ({columns}) VALUES ({placeholders})"
            )
            cur.execute(insert_query, form_data)
            conn.commit()
            cur.close()
            conn.close()

            QMessageBox.information(self, "Başarılı", "Form başarıyla gönderildi!")

            # İşlem tamamlandığında başarı mesajını göster
            self.show_success_message()

        except Exception as e:
            print(e)
            QMessageBox.critical(
                self, "Hata", f"E-posta gönderimi sırasında bir hata oluştu: {e}"
            )

    def show_success_message(self):
        # Tüm form alanlarını kaldır
        for i in reversed(range(self.form_layout.count())):
            widget = self.form_layout.itemAt(i).widget()
            if widget is not None:
                widget.deleteLater()

        # Teşekkür mesajını göster
        success_message = QLabel(
            "TEBRİKLER, BAŞARIYLA KAYIT OLDUNUZ.\nVIT PROJESİNE GÖSTERDİĞİNİZ İLGİYE TEŞEKKÜR EDERİZ."
        )
        success_message.setAlignment(Qt.AlignmentFlag.AlignCenter)
        success_message.setStyleSheet(
            "font-size: 18px; font-weight: bold; text-align: center;"
        )
        self.form_layout.addWidget(success_message)

        # Kapat butonunu ekle
        close_button = QPushButton("FORMU KAPAT")
        close_button.clicked.connect(self.close)
        self.form_layout.addWidget(close_button)


if __name__ == "__main__":
    app = QApplication(sys.argv)
    vit_form = VITForm()
    vit_form.show()
    sys.exit(app.exec())
